"""Small, source-image-coordinate helpers for building an editable PPTX.

The helper deliberately keeps the generated deck simple: one slide, explicit
surface properties on shapes created through this API, and warning-only text
fit estimates.  It never rewrites copy or silently shrinks text.
"""

from __future__ import annotations

import io
import json
import os
import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence

from PIL import Image, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Pt


PX_PER_INCH = 96.0
EMU_PER_INCH = 914400
# The bundled renderer is intentionally resolved relative to this module.
OFFICE_RENDERER = Path(__file__).with_name("render-powerpoint.ps1")


def _px(value: float) -> int:
    return int(round(float(value) * EMU_PER_INCH / PX_PER_INCH))


def _pt(value: float) -> float:
    return float(value) * 72.0 / PX_PER_INCH


def _rgb(value: Any) -> RGBColor:
    if isinstance(value, RGBColor):
        return value
    if isinstance(value, int):
        return RGBColor((value >> 16) & 255, (value >> 8) & 255, value & 255)
    if isinstance(value, (tuple, list)) and len(value) == 3:
        return RGBColor(*(int(part) for part in value))
    text = str(value).strip().lstrip("#")
    if len(text) != 6:
        raise ValueError(f"color must be #RRGGBB, RGB tuple, or integer: {value!r}")
    return RGBColor(int(text[0:2], 16), int(text[2:4], 16), int(text[4:6], 16))


def _is_none_color(value: Any) -> bool:
    return value is None or (
        isinstance(value, str) and value.strip().lower() in {"none", "transparent"}
    )


def _font_for_estimate(
    font_name: str,
    size_px: float,
    font_path: str | Path | None = None,
) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    """Load a local font for a warning-only width estimate.

    ``font_path`` is an optional measurement hint.  It never changes the font
    written into the PowerPoint file; the caller's ``font``/run font does that.
    """

    candidates: list[str] = []
    if font_path is not None:
        candidates.append(str(font_path))
    normalized = font_name.strip().lower().replace("_", " ")
    if os.name == "nt":
        if "yahei" in normalized or "microsoft yahei" in normalized or "msyh" in normalized:
            candidates.append(r"C:\Windows\Fonts\msyh.ttc")
        elif "simhei" in normalized or "黑体" in normalized:
            candidates.append(r"C:\Windows\Fonts\simhei.ttf")
        elif "simsun" in normalized or "宋体" in normalized:
            candidates.append(r"C:\Windows\Fonts\simsun.ttc")
    else:
        mac_fonts = [
            "/System/Library/Fonts/PingFang.ttc",
            "/System/Library/Fonts/STHeiti Light.ttc",
            "/Library/Fonts/Arial Unicode.ttf",
            "/System/Library/Fonts/Supplemental/Arial Unicode.ttf",
        ]
        if any(
            token in normalized
            for token in ("pingfang", "yahei", "microsoft yahei", "heiti", "黑体", "song", "宋体", "chinese")
        ):
            candidates.extend(mac_fonts)
        else:
            candidates.extend(
                [
                    "/System/Library/Fonts/Supplemental/Arial.ttf",
                    "/Library/Fonts/Arial.ttf",
                    *mac_fonts,
                ]
            )
    # Pillow accepts a font file path here.  A font family name is retained as
    # a best-effort candidate for environments that resolve it locally.
    candidates.extend([font_name, "Arial"])
    for candidate in candidates:
        try:
            return ImageFont.truetype(candidate, max(1, int(float(size_px))))
        except (OSError, ValueError):
            continue
    return ImageFont.load_default()


def _font_width(font: ImageFont.FreeTypeFont | ImageFont.ImageFont, text: str) -> float:
    if not text:
        return 0.0
    try:
        return float(font.getlength(text))
    except AttributeError:
        left, _top, right, _bottom = font.getbbox(text)
        return float(right - left)


def _suppress_inherited_effects(shape: Any) -> int:
    """Neutralize a theme effect reference on a newly created shape.

    Direct ``effectLst``/``effectDag`` children are explicit effects and are
    deliberately preserved.  Fill/line style references are also preserved;
    the helper's explicit fill/line properties take precedence for its shapes.
    """

    element = getattr(shape, "_element", None)
    if element is None:
        return 0
    style = element.find(qn("p:style"))
    if style is None:
        return 0
    effect_refs = list(style.findall(qn("a:effectRef")))
    if not effect_refs:
        return 0
    sp_pr = element.find(qn("p:spPr"))
    explicit_effect = sp_pr is not None and any(
        child.tag in {qn("a:effectLst"), qn("a:effectDag")} for child in sp_pr
    )
    if explicit_effect:
        return 0
    # CT_ShapeStyle requires effectRef when p:style exists. Removing just that
    # child creates a file PowerPoint rejects, even if python-pptx can read it.
    for effect_ref in effect_refs:
        effect_ref.set("idx", "0")
    if sp_pr is not None:
        sp_pr.append(OxmlElement("a:effectLst"))
    return len(effect_refs)


def _set_shape_surface(shape: Any, fill: Any = None, line: Any = None) -> None:
    """Set fill/line explicitly without erasing explicit effects."""

    _suppress_inherited_effects(shape)
    if hasattr(shape, "fill"):
        if _is_none_color(fill):
            shape.fill.background()
        else:
            shape.fill.solid()
            shape.fill.fore_color.rgb = _rgb(fill)
    if _is_none_color(line):
        shape.line.fill.background()
    else:
        shape.line.fill.solid()
        shape.line.color.rgb = _rgb(line)


def _set_line_options(shape: Any, width: float, dash: Any = None, arrow: Any = None) -> None:
    shape.line.width = Pt(_pt(width))
    ln = shape.line._get_or_add_ln()
    dash_value = {True: "dash", False: "solid", "dashed": "dash", "dash": "dash"}.get(dash, dash)
    dash_node = ln.find(qn("a:prstDash"))
    if dash_value:
        if dash_node is None:
            dash_node = ln.makeelement(qn("a:prstDash"), {})
            ln.append(dash_node)
        dash_node.set("val", str(dash_value))
    elif dash_node is not None:
        ln.remove(dash_node)

    for tag in ("a:headEnd", "a:tailEnd"):
        old = ln.find(qn(tag))
        if old is not None:
            ln.remove(old)
    arrow_text = str(arrow).lower() if arrow is not None else ""
    arrow_kind = "triangle"
    if arrow_text not in {"", "false", "none", "0"}:
        if arrow_text not in {"true", "1", "end", "start", "both"}:
            arrow_kind = arrow_text
        if arrow_text in {"true", "1", "start", "both"}:
            head = ln.makeelement(qn("a:headEnd"), {"type": arrow_kind, "w": "med", "len": "med"})
            ln.append(head)
        if arrow_text in {"true", "1", "end", "both"} or arrow_text not in {"start", "both"}:
            tail = ln.makeelement(qn("a:tailEnd"), {"type": arrow_kind, "w": "med", "len": "med"})
            ln.append(tail)


def _alignment(value: Any) -> PP_ALIGN:
    if isinstance(value, PP_ALIGN):
        return value
    return {
        "left": PP_ALIGN.LEFT,
        "center": PP_ALIGN.CENTER,
        "right": PP_ALIGN.RIGHT,
        "justify": PP_ALIGN.JUSTIFY,
    }.get(str(value).lower(), PP_ALIGN.LEFT)


def _bbox(values: Sequence[float], name: str) -> tuple[float, float, float, float]:
    if len(values) != 4:
        raise ValueError(f"{name} must be (x, y, w, h) in source-image pixels")
    x, y, w, h = (float(item) for item in values)
    if w <= 0 or h <= 0:
        raise ValueError(f"{name} width and height must be positive")
    return x, y, w, h


def _iter_shapes(shapes: Any) -> Iterable[Any]:
    """Yield every shape, including group children, exactly once."""

    for shape in shapes:
        yield shape
        children = getattr(shape, "shapes", None)
        if children is not None:
            yield from _iter_shapes(children)


def _shape_type(shape: Any) -> int | None:
    try:
        return int(shape.shape_type)
    except (AttributeError, TypeError, ValueError):
        return None


def _has_text(shape: Any) -> bool:
    try:
        return bool(shape.has_text_frame) and bool(str(shape.text).strip())
    except (AttributeError, ValueError):
        return False


def _render_soffice_preview(pptx_path: Path, preview_path: Path) -> dict[str, Any]:
    """Rasterize with LibreOffice when Windows PowerPoint is unavailable.

    Status is never ``ok``; callers must not treat this as Office verification.
    """

    renderer: dict[str, Any] = {
        "script": "soffice",
        "status": "unavailable",
        "office_verified": False,
    }
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    pdftoppm = shutil.which("pdftoppm")
    if not soffice:
        renderer["message"] = "LibreOffice (soffice) was not found for unofficial preview."
        return renderer
    if not pdftoppm:
        renderer["message"] = "pdftoppm was not found; unofficial LibreOffice preview skipped."
        return renderer
    with tempfile.TemporaryDirectory(prefix=".soffice-") as render_root:
        render_dir = Path(render_root)
        try:
            completed = subprocess.run(
                [
                    soffice,
                    "--headless",
                    "--convert-to",
                    "pdf",
                    "--outdir",
                    str(render_dir),
                    str(pptx_path),
                ],
                check=False,
                capture_output=True,
                text=True,
                timeout=180,
            )
        except (subprocess.TimeoutExpired, OSError) as exc:
            renderer.update(status="failed", message=f"LibreOffice preview failed: {exc}")
            return renderer
        if completed.returncode != 0:
            renderer.update(
                status="failed",
                message="LibreOffice returned a non-zero exit code.",
                stderr=(completed.stderr or "")[:400],
            )
            return renderer
        pdf = next(render_dir.glob("*.pdf"), None)
        if pdf is None:
            renderer.update(status="failed", message="LibreOffice produced no PDF.")
            return renderer
        prefix = render_dir / "page"
        raster = subprocess.run(
            ["pdftoppm", "-png", "-r", "144", "-f", "1", "-l", "1", str(pdf), str(prefix)],
            check=False,
            capture_output=True,
            text=True,
            timeout=60,
        )
        if raster.returncode != 0:
            renderer.update(
                status="failed",
                message="pdftoppm failed to rasterize the unofficial preview.",
            )
            return renderer
        rendered = next(render_dir.glob("page*.png"), None)
        if rendered is None:
            renderer.update(status="failed", message="pdftoppm produced no PNG.")
            return renderer
        shutil.copyfile(rendered, preview_path)
        renderer.update(
            status="unofficial-preview",
            message="Preview is LibreOffice/PDF raster, not Windows PowerPoint verification.",
        )
    return renderer


class Deck:
    """One-slide deck whose public geometry uses source PNG pixels.

    Coordinates stay in source pixels. When attached to the research library
    template, they are uniformly scaled into the 短科研图 / 长科研图 frame.
    """

    def __init__(
        self,
        source_png: str | Path,
        out_dir: str | Path,
        *,
        presentation: Presentation | None = None,
        slide: Any = None,
        origin_x: float = 0.0,
        origin_y: float = 0.0,
        scale: float | None = None,
        collection_path: str | Path | None = None,
        single_pptx_path: str | Path | None = None,
        slot: str | None = None,
    ):
        self.source_png = Path(source_png)
        self.out_dir = Path(out_dir)
        with Image.open(self.source_png) as image:
            self.width_px, self.height_px = image.size
        self.origin_x = float(origin_x)
        self.origin_y = float(origin_y)
        self.scale = float(EMU_PER_INCH / PX_PER_INCH if scale is None else scale)
        self.collection_path = Path(collection_path) if collection_path else None
        self.single_pptx_path = Path(single_pptx_path) if single_pptx_path else None
        self.slot = slot
        self._collection_slide_index: int | None = None
        if presentation is None:
            self.presentation = Presentation()
            self.presentation.slide_width = _px(self.width_px)
            self.presentation.slide_height = _px(self.height_px)
            self.slide = self.presentation.slides.add_slide(self.presentation.slide_layouts[6])
        else:
            self.presentation = presentation
            self.slide = slide if slide is not None else presentation.slides[-1]
        self.object_count = 0
        self.text_count = 0
        self.shape_count = 0
        self.image_count = 0
        self._text_warnings: list[str] = []

    @classmethod
    def into_library(
        cls,
        source_png: str | Path,
        slug: str,
        *,
        title: str,
        slot: str = "auto",
        out_dir: str | Path | None = None,
    ) -> "Deck":
        from library import (
            COLLECTION,
            LIBRARY_DIR,
            add_library_slide,
            archive_original,
            choose_slot,
            fit_transform,
        )

        archived = archive_original(source_png, slug)
        with Image.open(archived) as image:
            src_w, src_h = image.size
        chosen = choose_slot(src_w, src_h, slot)  # type: ignore[arg-type]
        origin_x, origin_y, scale, *_ = fit_transform(src_w, src_h, chosen)
        work = Path(out_dir) if out_dir else LIBRARY_DIR / f".build-{slug}"
        work.mkdir(parents=True, exist_ok=True)
        prs, slide, index = add_library_slide(COLLECTION, chosen, title)
        deck = cls(
            archived,
            work,
            presentation=prs,
            slide=slide,
            origin_x=origin_x,
            origin_y=origin_y,
            scale=scale,
            collection_path=COLLECTION,
            single_pptx_path=LIBRARY_DIR / f"{slug}.pptx",
            slot=chosen,
        )
        deck._collection_slide_index = index
        return deck

    def emu_x(self, x: float) -> int:
        return int(round(self.origin_x + float(x) * self.scale))

    def emu_y(self, y: float) -> int:
        return int(round(self.origin_y + float(y) * self.scale))

    def emu_len(self, value: float) -> int:
        return max(1, int(round(float(value) * self.scale)))

    def pt_size(self, size_px: float) -> float:
        return float(size_px) * self.scale * 72.0 / EMU_PER_INCH

    def _register(self, kind: str) -> None:
        # These counters provide useful live hints.  ``finish`` recomputes all
        # counts from the actual slide tree so direct additions are included.
        self.object_count += 1
        if kind == "text":
            self.text_count += 1
        elif kind == "image":
            self.image_count += 1
        else:
            self.shape_count += 1

    def text(
        self,
        x: float,
        y: float,
        w: float,
        h: float,
        text: str | Sequence[Mapping[str, Any]],
        size_px: float = 28,
        color: Any = "000000",
        bold: bool = False,
        align: Any = "left",
        font: str = "Arial",
        line_spacing: float = 1.0,
        font_path: str | Path | None = None,
        container_bbox: Sequence[float] | None = None,
    ) -> Any:
        """Add text with explicit run fonts and warning-only fit estimates.

        ``container_bbox`` checks the text-box rectangle against a declared
        container only.  It is not glyph-accurate and does not change the
        text, font size, or box geometry.
        """

        x, y, w, h = _bbox((x, y, w, h), "text bbox")
        container = _bbox(container_bbox, "container_bbox") if container_bbox is not None else None
        shape = self.slide.shapes.add_textbox(
            self.emu_x(x), self.emu_y(y), self.emu_len(w), self.emu_len(h)
        )
        _set_shape_surface(shape)
        frame = shape.text_frame
        frame.clear()
        frame.margin_left = frame.margin_right = 0
        frame.margin_top = frame.margin_bottom = 0
        frame.vertical_anchor = MSO_ANCHOR.TOP
        frame.word_wrap = True
        paragraph = frame.paragraphs[0]
        paragraph.alignment = _alignment(align)
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(0)
        paragraph.line_spacing = Pt(self.pt_size(size_px) * float(line_spacing))

        runs: Iterable[Mapping[str, Any]]
        if isinstance(text, str):
            runs = ({"text": text},)
        else:
            runs = text
        # Accumulate widths on shared visual lines.  A run after a newline
        # contributes to the current line rather than becoming a new max run.
        line_widths = [0.0]
        for run_spec in runs:
            run_text = str(run_spec.get("text", ""))
            run_font = str(run_spec.get("font", font))
            run_size = float(run_spec.get("size_px", size_px))
            run_bold = bool(run_spec.get("bold", bold))
            run_color = _rgb(run_spec.get("color", color))
            run_font_path = run_spec.get("font_path", font_path)
            font_file = _font_for_estimate(run_font, run_size, run_font_path)
            parts = run_text.split("\n")
            for part_index, part in enumerate(parts):
                if part:
                    run = paragraph.add_run()
                    run.text = part
                    run.font.name = run_font
                    run.font.size = Pt(self.pt_size(run_size))
                    run.font.bold = run_bold
                    run.font.color.rgb = run_color
                    r_pr = run._r.get_or_add_rPr()
                    for font_tag in ("a:latin", "a:ea"):
                        font_node = r_pr.find(qn(font_tag))
                        if font_node is None:
                            font_node = OxmlElement(font_tag)
                            r_pr.append(font_node)
                        font_node.set("typeface", run_font)
                line_widths[-1] += _font_width(font_file, part)
                if part_index < len(parts) - 1:
                    paragraph._p.append(OxmlElement("a:br"))
                    line_widths.append(0.0)
        estimate_width = max(line_widths, default=0.0)
        estimate_lines = max(1, len(line_widths))
        estimate_height = estimate_lines * float(size_px) * 1.2 * float(line_spacing)
        if estimate_width > w or estimate_height > h:
            self._text_warnings.append(
                f"text box {self.text_count + 1} may need visual review "
                f"(estimated {estimate_width:.0f}x{estimate_height:.0f}px in {w:.0f}x{h:.0f}px)"
            )
        if container is not None:
            cx, cy, cw, ch = container
            if x < cx or y < cy or x + w > cx + cw or y + h > cy + ch:
                self._text_warnings.append(
                    f"text box {self.text_count + 1} extends beyond container_bbox; "
                    "this check compares text-box bounds only and is not glyph-accurate"
                )
        self._register("text")
        return shape

    def rect(
        self,
        x: float,
        y: float,
        w: float,
        h: float,
        fill: Any = None,
        line: Any = None,
        radius: Any = False,
    ) -> Any:
        x, y, w, h = _bbox((x, y, w, h), "rect bbox")
        shape_type = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
        shape = self.slide.shapes.add_shape(
            shape_type, self.emu_x(x), self.emu_y(y), self.emu_len(w), self.emu_len(h)
        )
        _set_shape_surface(shape, fill, line)
        self._register("shape")
        return shape

    def line(
        self,
        x1: float,
        y1: float,
        x2: float,
        y2: float,
        color: Any = "000000",
        width: float = 1,
        dash: Any = None,
        arrow: Any = None,
    ) -> Any:
        shape = self.slide.shapes.add_connector(
            1, self.emu_x(x1), self.emu_y(y1), self.emu_x(x2), self.emu_y(y2)
        )
        _set_shape_surface(shape, None, color)
        _set_line_options(shape, width * self.scale * PX_PER_INCH / EMU_PER_INCH, dash, arrow)
        self._register("shape")
        return shape

    def crop(
        self,
        source_bbox: Sequence[float],
        target_bbox: Sequence[float] | None = None,
    ) -> Any:
        """Place a cropped source PNG region; boxes are ``(x, y, w, h)`` pixels."""

        sx, sy, sw, sh = _bbox(source_bbox, "source_bbox")
        tx, ty, tw, th = _bbox(target_bbox or source_bbox, "target_bbox")
        if sx < 0 or sy < 0 or sx + sw > self.width_px or sy + sh > self.height_px:
            raise ValueError("source_bbox must be inside the source PNG")
        left, top = int(round(sx)), int(round(sy))
        right, bottom = int(round(sx + sw)), int(round(sy + sh))
        if left < 0 or top < 0 or right > self.width_px or bottom > self.height_px:
            raise ValueError("source_bbox must be inside the source PNG")
        if right <= left or bottom <= top:
            raise ValueError("source_bbox must cover at least one source pixel")
        with Image.open(self.source_png) as image:
            crop_image = image.convert("RGBA").crop((left, top, right, bottom))
            stream = io.BytesIO()
            crop_image.save(stream, format="PNG")
        stream.seek(0)
        picture = self.slide.shapes.add_picture(
            stream, self.emu_x(tx), self.emu_y(ty), self.emu_len(tw), self.emu_len(th)
        )
        self._register("image")
        return picture

    def normalize_inherited_effects(self, shapes: Any = None) -> int:
        """Explicitly neutralize inherited theme effects on direct shapes.

        This is opt-in because the caller may be working with template or
        manually styled objects.  Explicit ``effectLst``/``effectDag`` nodes
        are preserved.  With no argument, the current slide tree is visited.
        """

        if shapes is None:
            roots = self.slide.shapes
        elif hasattr(shapes, "_element"):
            roots = (shapes,)
        else:
            roots = shapes
        changed = 0
        for shape in _iter_shapes(roots):
            changed += _suppress_inherited_effects(shape)
        return changed

    def _recount(self) -> None:
        shapes = list(_iter_shapes(self.slide.shapes))
        self.object_count = len(shapes)
        self.text_count = sum(1 for shape in shapes if _has_text(shape))
        self.image_count = sum(1 for shape in shapes if _shape_type(shape) == 13)
        self.shape_count = self.object_count - self.text_count - self.image_count

    def _save_report(self, report_path: Path, renderer: Mapping[str, Any]) -> None:
        report = {
            "source_size_px": [self.width_px, self.height_px],
            "slot": self.slot,
            "scale_emu_per_px": self.scale,
            "object_count": self.object_count,
            "text_count": self.text_count,
            "shape_count": self.shape_count,
            "image_count": self.image_count,
            "text_fit_warnings": list(self._text_warnings),
            "renderer": dict(renderer),
            "render_validation": dict(renderer),
        }
        report_path.write_text(
            json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
        )

    def _render(
        self,
        pptx_path: Path,
        preview_path: Path,
        verification_path: Path,
    ) -> dict[str, Any]:
        renderer: dict[str, Any] = {"script": Path(OFFICE_RENDERER).name, "status": "unavailable"}
        if os.name != "nt":
            unofficial = _render_soffice_preview(pptx_path, preview_path)
            if unofficial.get("status") == "unofficial-preview":
                return unofficial
            renderer["message"] = (
                "Native Windows PowerPoint rendering is unavailable outside Windows. "
                + unofficial.get("message", "")
            ).strip()
            return renderer
        renderer_script = Path(OFFICE_RENDERER)
        if not renderer_script.exists():
            renderer["message"] = "Bundled native Windows Office renderer was not found."
            return renderer
        powershell = shutil.which("powershell.exe") or shutil.which("powershell")
        if not powershell:
            renderer["message"] = "PowerShell is unavailable; native Windows Office rendering was not run."
            return renderer

        with tempfile.TemporaryDirectory(prefix=".render-", dir=self.out_dir) as render_root:
            render_dir = Path(render_root)
            command = [
                powershell,
                "-NoProfile",
                "-ExecutionPolicy",
                "Bypass",
                "-File",
                str(renderer_script),
                "-DeckPath",
                str(pptx_path),
                "-OutputDirectory",
                str(render_dir),
            ]
            try:
                completed = subprocess.run(
                    command,
                    check=False,
                    capture_output=True,
                    text=True,
                    timeout=180,
                )
            except subprocess.TimeoutExpired as exc:
                renderer.update(
                    status="failed",
                    message="Native Windows Office renderer timed out.",
                    exit_code=None,
                )
                if exc.stderr:
                    renderer["stderr"] = str(exc.stderr)[:400]
                return renderer
            except OSError as exc:
                renderer.update(status="failed", message=f"Could not start native renderer: {exc}")
                return renderer

            renderer["exit_code"] = completed.returncode
            if completed.stderr.strip():
                renderer["stderr"] = completed.stderr.strip()[:400]
            if completed.returncode != 0:
                renderer.update(
                    status="failed",
                    message="Native Windows Office renderer returned a non-zero exit code.",
                )
                return renderer
            rendered = next(render_dir.rglob("page-01.png"), None)
            verification = next(render_dir.rglob("powerpoint-verification.json"), None)
            if verification is not None:
                shutil.copyfile(verification, verification_path)
                renderer["verification"] = str(verification_path)
            if rendered is None:
                renderer.update(
                    status="failed",
                    message="Native renderer succeeded but produced no page-01.png.",
                )
                return renderer
            shutil.copyfile(rendered, preview_path)
            if verification is None:
                renderer.update(
                    status="missing-verification",
                    message="Native renderer produced a preview without verification JSON.",
                )
            else:
                renderer["status"] = "ok"
        return renderer

    def finish(self, render: bool = True) -> dict[str, Any]:
        """Save deliverables and optionally render with native Windows Office.

        Rendering is best-effort and its status is explicit in the returned
        mapping and editability report.  Any old preview or verification file
        is removed before the new attempt, so a failed attempt cannot expose a
        stale preview as current evidence.
        """

        self._recount()
        deliverables = self.out_dir / "deliverables"
        deliverables.mkdir(parents=True, exist_ok=True)
        pptx_path = deliverables / "editable.pptx"
        report_path = deliverables / "editability-report.json"
        preview_path = deliverables / "preview.png"
        verification_path = deliverables / "powerpoint-verification.json"
        for stale in (preview_path, verification_path):
            if stale.exists():
                stale.unlink()
        self.presentation.save(pptx_path)
        if self.collection_path is not None:
            self.presentation.save(self.collection_path)
            if self.single_pptx_path is not None and self._collection_slide_index is not None:
                from library import save_single_slide

                save_single_slide(
                    self.collection_path,
                    self._collection_slide_index,
                    self.single_pptx_path,
                )

        if render:
            renderer = self._render(pptx_path, preview_path, verification_path)
        else:
            renderer = {
                "script": Path(OFFICE_RENDERER).name,
                "status": "skipped",
                "message": "Rendering disabled by finish(render=False).",
            }
        renderer["preview"] = str(preview_path) if preview_path.exists() else None
        self._save_report(report_path, renderer)
        preview_display = str(preview_path) if preview_path.exists() else "<missing>"
        print(
            f"preview={preview_display} objects={self.object_count} "
            f"text_warnings={len(self._text_warnings)} render_status={renderer['status']}"
        )
        return {
            "pptx": str(self.single_pptx_path or pptx_path),
            "collection": str(self.collection_path) if self.collection_path else None,
            "original": str(self.source_png),
            "slot": self.slot,
            "report": str(report_path),
            "preview": str(preview_path) if preview_path.exists() else None,
            "objects": self.object_count,
            "text_warnings": list(self._text_warnings),
            "renderer": renderer,
            "rendered": renderer["status"] == "ok",
        }


__all__ = ["Deck", "OFFICE_RENDERER"]
