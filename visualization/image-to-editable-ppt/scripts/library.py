"""科研图表库：原图归档 + 按模板短图/长图边界等比放入可编辑 PPT。"""

from __future__ import annotations

import copy
import shutil
import os
from pathlib import Path
from typing import Any, Literal

from pptx import Presentation
from pptx.oxml.ns import qn

LIBRARY_DIR = Path(
    os.environ.get("XY_VIZ_LIBRARY")
    or os.environ.get("IMAGE_TO_PPT_LIBRARY")
    or (Path.home() / "Documents" / "科研参考图")
).expanduser()
COLLECTION = LIBRARY_DIR / os.environ.get("XY_VIZ_COLLECTION", "科研图表集合.pptx")

# 从 科研图表集合.pptx 模版图页读出的图框（EMU）。短图留右侧栏，长图拉满内容区。
SHORT_BOX = (554610, 1351499, 8955024, 4918672)
LONG_BOX = (554610, 1351499, 11082654, 4918672)
INSET_EMU = 45720  # 0.05" inside the black frame
TEMPLATE_SLIDE_COUNT = 2  # slide 1 短科研图, slide 2 长科研图

Slot = Literal["short", "long", "auto"]


def box_for(slot: Literal["short", "long"]) -> tuple[int, int, int, int]:
    return SHORT_BOX if slot == "short" else LONG_BOX


def choose_slot(src_w: int, src_h: int, slot: Slot = "auto") -> Literal["short", "long"]:
    if slot in {"short", "long"}:
        return slot
    fill = {}
    for name in ("short", "long"):
        _origin_x, _origin_y, scale, used_w, used_h, box_w, box_h = fit_transform(
            src_w, src_h, name
        )
        fill[name] = (used_w * used_h) / (box_w * box_h)
    return "short" if fill["short"] >= fill["long"] else "long"


def fit_transform(
    src_w: int, src_h: int, slot: Literal["short", "long"]
) -> tuple[float, float, float, float, float, int, int]:
    """Uniform scale: source pixels → slide EMU, letterboxed inside the figure frame."""

    left, top, box_w, box_h = box_for(slot)
    inner_w = box_w - 2 * INSET_EMU
    inner_h = box_h - 2 * INSET_EMU
    scale = min(inner_w / float(src_w), inner_h / float(src_h))
    used_w = src_w * scale
    used_h = src_h * scale
    origin_x = left + INSET_EMU + (inner_w - used_w) / 2.0
    origin_y = top + INSET_EMU + (inner_h - used_h) / 2.0
    return origin_x, origin_y, scale, used_w, used_h, box_w, box_h


def archive_original(source_png: str | Path, slug: str) -> Path:
    LIBRARY_DIR.mkdir(parents=True, exist_ok=True)
    dest = LIBRARY_DIR / f"{slug}.png"
    src = Path(source_png)
    if src.resolve() != dest.resolve():
        shutil.copyfile(src, dest)
    return dest


def _duplicate_slide(prs: Presentation, index: int) -> Any:
    source = prs.slides[index]
    dest = prs.slides.add_slide(source.slide_layout)
    sp_tree = dest.shapes._spTree
    for child in list(sp_tree):
        tag = child.tag.split("}")[-1]
        if tag in {"sp", "pic", "cxnSp", "grpSp"}:
            sp_tree.remove(child)
    for child in source.shapes._spTree:
        tag = child.tag.split("}")[-1]
        if tag in {"sp", "pic", "cxnSp", "grpSp"}:
            sp_tree.append(copy.deepcopy(child))
    return dest


def _is_figure_frame(shape: Any, slot: Literal["short", "long"]) -> bool:
    _l, _t, w, h = box_for(slot)
    try:
        return abs(int(shape.width) - w) < 2000 and abs(int(shape.height) - h) < 2000
    except Exception:
        return False


def _clear_frame_label(slide: Any, slot: Literal["short", "long"]) -> None:
    for shape in slide.shapes:
        if not _is_figure_frame(shape, slot):
            continue
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    run.text = ""
        try:
            from pptx.dml.color import RGBColor

            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
        except Exception:
            pass


def _set_title(slide: Any, title: str) -> None:
    for shape in slide.shapes:
        try:
            if shape.is_placeholder and "TITLE" in str(shape.placeholder_format.type):
                shape.text = title
                return
        except Exception:
            continue
    for shape in slide.shapes:
        if shape.has_text_frame and str(shape.name).startswith("标题") and int(shape.top or 0) < 1_000_000:
            shape.text = title
            return


def add_library_slide(
    collection_path: Path | str,
    slot: Literal["short", "long"],
    title: str,
) -> tuple[Presentation, Any, int]:
    prs = Presentation(str(collection_path))
    index = 0 if slot == "short" else 1
    slide = _duplicate_slide(prs, index)
    _clear_frame_label(slide, slot)
    if title:
        _set_title(slide, title)
    return prs, slide, len(prs.slides) - 1


def save_single_slide(collection_path: Path, slide_index: int, dest: Path) -> None:
    shutil.copyfile(collection_path, dest)
    prs = Presentation(str(dest))
    sld_id_lst = prs.slides._sldIdLst
    keep = sld_id_lst[slide_index]
    for el in list(sld_id_lst):
        if el is keep:
            continue
        r_id = el.get(qn("r:id"))
        if r_id:
            try:
                prs.part.drop_rel(r_id)
            except Exception:
                pass
        sld_id_lst.remove(el)
    prs.save(str(dest))
